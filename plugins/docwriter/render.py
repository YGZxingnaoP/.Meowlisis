# -*- coding: utf-8 -*-
"""docwriter 渲染层 v2：约束生成 + 确定性排版，一次成型（不做事后重排）。

* DOCX：样式表 + 封面 + 页眉页脚 + 目录域 + 图表自动编号 + 三线表 + 圆角插图 + 章首分页
* PPTX：版式函数库（layouts）自绘版面；模型只选版式与填内容
* 内容归一：条数/字数硬上限在此兜底（确定性截断，非迭代）

对外接口（向后兼容）
--------------------
    write_doc(filename, content, output_dir, **opts) -> path
    write_ppt(filename, slides, output_dir, **opts) -> path
"""

import os
import re

IMAGE_DIR = "assets"
TEXT_AREA_CM = 15.2          # A4 默认页边距下的可用宽度

# 内容契约（渲染层兜底，确定性截断；提示词里同步约束模型）
LIMIT = {
    "bullets": 5, "toc": 7, "kpi": 3, "timeline": 5, "matrix": 4, "grid": 4,
    "two_col": 4, "points": 4, "table_rows": 8, "table_cols": 4,
}
CHARS = {"bullets": 40, "toc": 20, "points": 40, "caption": 30, "cell": 18, "matrix": 40}


# ==================== 通用 ====================

def _resolve(output_dir, sub=""):
    """解析并创建输出目录"""
    if not os.path.isabs(output_dir):
        output_dir = os.path.join(os.getcwd(), output_dir)
    if sub:
        output_dir = os.path.join(output_dir, sub)
    os.makedirs(output_dir, exist_ok=True)
    return output_dir


def _filename(filename, ext):
    """生成安全文件名"""
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", str(filename or "").strip())
    name = os.path.splitext(name)[0].strip(" ._")[:60]
    return f"{name or '文档'}.{ext}"


def _unique_path(folder, filename):
    """生成不冲突的完整路径"""
    base, ext = os.path.splitext(filename)
    path = os.path.join(folder, filename)
    index = 2
    while os.path.exists(path):
        path = os.path.join(folder, f"{base}({index}){ext}")
        index += 1
    return path


def resolve_asset(path, base_dir):
    """把图片相对路径解析成存在的绝对路径（支持 assets/ 前缀与反斜杠）"""
    p = str(path or "").strip().strip('"\'')
    if not p:
        return ""
    p = p.replace("\\", "/")
    if os.path.isabs(p) and os.path.isfile(p):
        return p
    for cand in (os.path.join(base_dir, p),
                 os.path.join(base_dir, IMAGE_DIR, os.path.basename(p)),
                 os.path.join(os.getcwd(), p),
                 p):
        if cand and os.path.isfile(cand):
            return os.path.abspath(cand)
    return ""


def assets_dir(output_dir):
    """图片统一落盘目录（不存在则创建）"""
    return _resolve(output_dir, IMAGE_DIR)


def norm_items(value, limit=None, max_chars=None):
    """把要点归一化成列表：按行/分号拆、去空、限量、超长截断（确定性）"""
    if value is None:
        return []
    if isinstance(value, str):
        parts = re.split(r"[\n；;]+", value)
        out = [p.strip(" \t-•") for p in parts if p.strip(" \t-•")]
    elif isinstance(value, (list, tuple)):
        out = [str(v).strip() for v in value if str(v).strip()]
    else:
        out = [str(value)]
    if limit:
        out = out[:limit]
    if max_chars:
        out = [i if len(i) <= max_chars else i[:max_chars - 1] + "…" for i in out]
    return out


# ==================== DOCX ====================

_MD_IMAGE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
_MD_TABLE_SEP = re.compile(r"^\|[\s:|-]+\|$")


def _split_table(lines, i):
    """从 i 开始解析 Markdown 表格，返回 (表头, 数据行, 下一行索引)"""
    block = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        block.append(lines[i].strip())
        i += 1
    rows = [[c.strip() for c in ln.strip("|").split("|")] for ln in block]
    if len(rows) >= 2 and _MD_TABLE_SEP.match(block[1]):
        return rows[0], rows[2:], i
    return (rows[0] if rows else []), rows[1:], i


def _cell_border(cell, edge="bottom", color="BFBFBF", size="6"):
    """给单元格设置单边线（三线表用）"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:color"), str(color).lstrip("#").upper())
    borders.append(el)


def _docx_image(doc, path, caption="", theme=None, folder=""):
    """居中插图（圆角+描边）+ 图注"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm
    from PIL import Image

    src = path
    if theme is not None:
        src = D_mod().prepare_image(path, folder or os.path.dirname(path), theme)
    try:
        iw, ih = Image.open(src).size
    except Exception:
        return False
    w = TEXT_AREA_CM
    if iw and ih and ih > iw * 1.35:
        w = w * 0.62
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = 0
    try:
        p.add_run().add_picture(src, width=Cm(w))
    except Exception:
        return False
    if caption:
        c = doc.add_paragraph(str(caption), style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return True


def _docx_table(doc, header, rows, theme=None):
    """三线表：顶线 + 表头下线 + 底线（比全框线专业）"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    raw_cols = max(len(header), max((len(r) for r in rows), default=0)) or 1
    cols = min(raw_cols, LIMIT["table_cols"])
    header = (list(header) + [""] * cols)[:cols]
    rows = [(list(r) + [""] * cols)[:cols] for r in rows[:LIMIT["table_rows"]]]
    tbl = doc.add_table(rows=1 + len(rows), cols=cols)
    line = "BFBFBF"
    for ci, val in enumerate(header):
        cell = tbl.cell(0, ci)
        cell.text = str(val)
        _cell_border(cell, "top", line, "12")
        _cell_border(cell, "bottom", line, "8")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = 0
            for r in p.runs:
                r.font.bold = True
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = tbl.cell(ri + 1, ci)
            cell.text = str(row[ci]) if ci < len(row) else ""
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = 0
            if ri == len(rows) - 1:
                _cell_border(cell, "bottom", line, "12")
    doc.add_paragraph()


def D_mod():
    """延迟取设计模块（避免顶层循环导入）"""
    from . import design

    return design


def write_doc(filename, content, output_dir, theme=None, title=None, subtitle="",
              meta=None, cover=True, toc=False):
    """生成 Word 文档，返回路径

    content 为简化 Markdown：
        # / ## / ### 标题、- 项目符号、1. 编号、> 引用、
        | 表格 |、![图注](图片路径)、--- 分页
    图表编号由渲染层自动生成（图 1 / 表 1），正文不必手写编号。
    """
    from docx import Document

    D = D_mod()
    th = D.get_theme(theme)
    folder = _resolve(output_dir)
    path = _unique_path(folder, _filename(filename, "docx"))
    doc = Document()
    D.apply_docx_theme(doc, th)
    doc_title = str(title or "").strip()
    D.docx_header(doc, th, doc_title)
    D.docx_page_footer(doc, th, text=doc_title)

    lines = str(content or "").splitlines()
    if cover:
        D.docx_cover(doc, th, doc_title or (lines[0].lstrip("# ").strip() if lines else "文档"),
                     subtitle, meta or [])
    if toc:
        D.docx_toc_field(doc)

    fig_no = tbl_no = 0
    written = 0
    first_heading_seen = False
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()
        i += 1
        if not s:
            continue

        m = _MD_IMAGE.match(s)
        if m:
            fig_no += 1
            cap = re.sub(r"^(图|图\s*\d+|Figure)[\s\.:：]*\d*[\s\.:：]*", "", m.group(1).strip())
            caption = f"图 {fig_no}  {cap}" if cap else f"图 {fig_no}"
            real = resolve_asset(m.group(2).strip(), folder)
            if real:
                _docx_image(doc, real, caption, th, folder)
            else:
                doc.add_paragraph(f"（配图缺失：{m.group(2).strip()}）", style="Caption")
            written += 1
            continue

        if s.startswith("|"):
            tbl_no += 1
            header, rows, i = _split_table(lines, i - 1)
            _docx_table(doc, header, rows, th)
            c = doc.add_paragraph(f"表 {tbl_no}", style="Caption")
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            written += 1
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            doc.add_page_break()
            continue

        if s.startswith("#### "):
            doc.add_heading(s[5:].strip(), level=3)
        elif s.startswith("### "):
            doc.add_heading(s[4:].strip(), level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=2)
        elif s.startswith("# "):
            head = s[2:].strip()
            if cover and head == doc_title:
                continue
            if first_heading_seen:
                doc.add_page_break()          # 章首另起一页
            first_heading_seen = True
            doc.add_heading(head, level=1)
        elif s.startswith("> "):
            p = doc.add_paragraph(s[2:].strip())
            try:
                p.style = doc.styles["Quote"]
            except Exception:
                pass
            p.paragraph_format.first_line_indent = 0
        elif s.startswith(("- ", "* ", "• ")):
            doc.add_paragraph(s[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", s):
            doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", s), style="List Number")
        else:
            p = doc.add_paragraph(s)
            if s.startswith("（") or s.startswith("注："):
                p.style = doc.styles["Caption"]
        written += 1

    if not written:
        doc.add_paragraph("（内容为空）")
    doc.save(path)
    return path


# ==================== PPTX ====================

def pick_layout(item, index, total):
    """未指定 layout 时按内容自动选版式"""
    from . import layouts as L

    layout = str(item.get("layout") or "").strip().lower()
    if layout in L.LAYOUTS:
        return layout
    if index == 0:
        return "cover"
    bullets = norm_items(item.get("bullets") or item.get("items"))
    has_media = bool(item.get("image") or item.get("images"))
    if index == total - 1 and not bullets and not has_media:
        return "end"
    if item.get("header") and item.get("rows"):
        return "table"
    if item.get("cards"):
        return "kpi"
    if item.get("nodes"):
        return "timeline"
    if item.get("cells"):
        return "matrix"
    if item.get("images"):
        return "grid"
    if item.get("image") and not bullets:
        return "full_image"
    if item.get("image") and norm_items(item.get("points")):
        return "chart"
    if item.get("image") and bullets:
        return "image_right"
    if item.get("left_items") or item.get("right_items"):
        return "two_col"
    if str(item.get("kind") or "").lower() == "quote" or item.get("author"):
        return "quote"
    return "bullets"


def _slides_images(value, folder):
    """归一化 grid 版式的图片列表"""
    out = []
    for it in (value or []):
        if isinstance(it, dict):
            out.append({"path": resolve_asset(it.get("path"), folder) or it.get("path") or "",
                        "caption": str(it.get("caption") or "")[:CHARS["caption"]]})
        else:
            out.append({"path": resolve_asset(it, folder) or str(it), "caption": ""})
    return out[:LIMIT["grid"]]


def _cards(value):
    """归一化 kpi 卡片"""
    out = []
    for c in (value or []):
        if isinstance(c, dict):
            out.append({"value": str(c.get("value") or c.get("num") or "—")[:10],
                        "label": str(c.get("label") or "")[:16],
                        "note": str(c.get("note") or "")[:36]})
        else:
            out.append({"value": str(c)[:10], "label": "", "note": ""})
    return out[:LIMIT["kpi"]]


def _nodes(value):
    """归一化 timeline 节点"""
    out = []
    for n in (value or []):
        if isinstance(n, dict):
            out.append({"label": str(n.get("label") or "")[:8],
                        "title": str(n.get("title") or "")[:10],
                        "desc": str(n.get("desc") or "")[:14]})
        else:
            out.append({"label": "", "title": str(n)[:10], "desc": ""})
    return out[:LIMIT["timeline"]]


def _cells(value):
    """归一化 matrix 单元格"""
    out = []
    for c in (value or []):
        if isinstance(c, dict):
            out.append({"title": str(c.get("title") or "")[:16],
                        "items": norm_items(c.get("items") or c.get("bullets") or c.get("text"),
                                            limit=3, max_chars=CHARS["matrix"])})
        else:
            out.append({"title": str(c)[:16], "items": []})
    return out[:LIMIT["matrix"]]


def write_ppt(filename, slides, output_dir, theme=None, title=None, subtitle="", meta=""):
    """生成 PPT，返回路径

    slides 元素字段：layout / title / kicker / bullets / points / image / caption /
        header+rows / cards / nodes / cells / images / text+author / no+desc /
        left_title+left_items+right_title+right_items
    条数与字数上限由渲染层兜底（bullets ≤5 条、toc ≤7 条等），保证版面不溢出。
    """
    from pptx import Presentation
    from pptx.util import Cm

    from . import design as D
    from . import layouts as L

    th = D.get_theme(theme)
    folder = _resolve(output_dir)
    path = _unique_path(folder, _filename(filename, "pptx"))
    prs = Presentation()
    prs.slide_width = Cm(D.PAGE["ppt_w"])
    prs.slide_height = Cm(D.PAGE["ppt_h"])

    if isinstance(slides, dict):
        slides = slides.get("slides") or []
    if not isinstance(slides, list) or not slides:
        slides = [{"title": str(title or filename or "演示文稿"), "bullets": []}]
    if title and not (slides and slides[0].get("title")):
        slides[0]["title"] = title

    doc_title = str(title or "").strip()
    total = len(slides)
    page = 0
    for idx, raw in enumerate(slides):
        item = raw if isinstance(raw, dict) else {"title": str(raw)}
        layout = pick_layout(item, idx, total)
        d = dict(item)
        img = resolve_asset(d.get("image"), folder) or d.get("image") or ""
        bullets = norm_items(d.get("bullets") or d.get("items"), LIMIT["bullets"], CHARS["bullets"])
        points = norm_items(d.get("points"), LIMIT["points"], CHARS["points"])
        page_no = None
        if layout not in ("cover", "end", "section", "quote"):
            page += 1
            page_no = d.get("page_no") or page
        common = {"page_no": page_no, "doc_title": doc_title}

        if layout == "cover":
            L.cover(prs, th, d.get("title") or doc_title or "演示文稿",
                    d.get("subtitle") or subtitle, d.get("meta") or meta)
        elif layout == "end":
            L.end(prs, th, d.get("title") or "谢谢观看", d.get("subtitle") or "")
        elif layout == "quote":
            L.quote(prs, th, d.get("text") or d.get("title") or "", d.get("author") or "",
                    page_no)
        elif layout == "kpi":
            L.kpi(prs, th, d.get("title"), _cards(d.get("cards")), d.get("kicker"), **common)
        elif layout == "timeline":
            L.timeline(prs, th, d.get("title"), _nodes(d.get("nodes")), **common)
        elif layout == "matrix":
            L.matrix(prs, th, d.get("title"), _cells(d.get("cells")), **common)
        elif layout == "grid":
            L.grid(prs, th, d.get("title"), _slides_images(d.get("images"), folder), **common)
        elif layout == "table":
            L.table(prs, th, d.get("title"), d.get("header"), d.get("rows"), **common)
        elif layout == "chart":
            L.chart(prs, th, d.get("title"), img, points, d.get("caption") or "", **common)
        elif layout == "full_image":
            L.full_image(prs, th, d.get("title"), img, d.get("caption") or "", points, **common)
        elif layout == "image_right":
            L.image_right(prs, th, d.get("title"), bullets, img, d.get("caption") or "", **common)
        elif layout == "two_col":
            L.two_col(prs, th, d.get("title"), d.get("left_title"),
                      norm_items(d.get("left_items"), LIMIT["two_col"], 24), d.get("right_title"),
                      norm_items(d.get("right_items"), LIMIT["two_col"], 24), **common)
        elif layout == "toc":
            L.toc(prs, th, norm_items(d.get("bullets") or d.get("toc"), LIMIT["toc"],
                                      CHARS["toc"]), **common)
        elif layout == "section":
            L.section(prs, th, d.get("title"), d.get("no") or "", d.get("desc") or "", **common)
        else:
            L.bullets(prs, th, d.get("title"), bullets, d.get("kicker"), **common)

    prs.save(path)
    return path
